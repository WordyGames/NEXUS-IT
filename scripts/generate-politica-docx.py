#!/usr/bin/env python3
"""Genera el Word de la Política de Uso y Cambio de Equipo de Cómputo."""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PRIMARY = RGBColor(0x15, 0x65, 0xC0)
DARK = RGBColor(0x1a, 0x1a, 0x1a)
GRAY = RGBColor(0x55, 0x55, 0x55)

OUT_PATH = Path(__file__).parent.parent / 'docs' / 'Politica_Uso_y_Cambio_Equipo_Computo.docx'


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def style_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
            if header and row_idx == 0:
                set_cell_shading(cell, '1565C0')
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, 'F2F6FC')


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = PRIMARY
    return h


def add_body(doc, text, bold=False, size=10.5, color=DARK, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        run = p.add_run(item)
        run.font.size = Pt(10.5)
        p.paragraph_format.space_after = Pt(4)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.font.bold = True
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = ''
            run = row_cells[i].paragraphs[0].add_run(val)
    style_table(table)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def build():
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    # ===== PORTADA =====
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(60)
    run = title.add_run('POLÍTICA DE USO Y CAMBIO\nDE EQUIPO DE CÓMPUTO')
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = PRIMARY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(20)
    run = subtitle.add_run('Especias Naturales del Norte  ·  Grupo AMEX  ·  Equipos Osenal (Liumaq)')
    run.font.size = Pt(13)
    run.font.color.rgb = GRAY

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(200)
    for line in [
        'Área responsable: Sistemas / TI (Nexus IT)',
        'Versión: 1.0',
        'Vigencia desde: ____________________',
    ]:
        r = meta.add_run(line + '\n')
        r.font.size = Pt(11)
        r.font.color.rgb = DARK

    doc.add_page_break()

    # ===== CONTENIDO =====
    add_heading(doc, '1. Objetivo', 1)
    add_body(doc, 'Establecer las reglas de uso, cuidado, mantenimiento y reemplazo de los '
                  'equipos de cómputo propiedad de la empresa, con el fin de proteger la '
                  'inversión en activos, garantizar la continuidad operativa y definir '
                  'criterios claros sobre cuándo un equipo se repara y cuándo se cambia.')

    add_heading(doc, '2. Alcance', 1)
    add_body(doc, 'Aplica a todos los equipos registrados en Nexus IT, sin importar la '
                  'empresa del grupo a la que pertenezcan. Todo equipo asignado a un '
                  'colaborador debe estar respaldado por una Carta Responsiva firmada '
                  '(empleado + RH/Sistemas) antes de su entrega.')

    add_heading(doc, '3. Definiciones', 1)
    add_table(doc, ['Término', 'Significado'], [
        ['Equipo', 'Cualquier dispositivo registrado en Nexus IT: desktop, laptop, '
                    'teléfono, tablet, servidor, impresora, router o switch.'],
        ['Carta Responsiva', 'Documento individual que formaliza la entrega de un equipo '
                              'a un colaborador y sus condiciones de uso.'],
        ['Ticket', 'Solicitud de soporte registrada en Nexus IT (categorías: hardware, '
                    'software, red, correo, impresión, accesos, otro).'],
        ['Estado del equipo', 'Activo, Inactivo, En mantenimiento o Baja, según el ciclo '
                               'de vida registrado en el sistema.'],
    ], widths=[1.7, 4.8])

    add_heading(doc, '4. Asignación de equipo', 1)
    add_numbered(doc, [
        'Todo equipo se entrega mediante Carta Responsiva generada desde Nexus IT, '
        'firmada por el colaborador y por RH/Sistemas.',
        'El equipo es propiedad de la empresa y se asigna únicamente para el desempeño '
        'de las funciones del puesto.',
        'La Carta Responsiva debe incluir: datos del colaborador, especificaciones del '
        'equipo (marca, modelo, número de serie, IMEI si aplica) y condiciones aceptadas.',
        'Cambios de puesto, área o empresa dentro del grupo requieren revisión de la '
        'asignación vigente.',
    ])

    add_heading(doc, '5. Uso aceptable', 1)
    add_numbered(doc, [
        'El equipo se destina exclusivamente a actividades laborales, salvo autorización '
        'expresa de Sistemas.',
        'Prohibido instalar software, extensiones o aplicaciones no autorizadas por '
        'Sistemas.',
        'Prohibido compartir contraseñas, cuentas corporativas (correo, Google, VPN) o '
        'prestar el equipo a terceros no autorizados.',
        'Prohibido modificar configuración de hardware o software (desinstalar antivirus, '
        'alterar BIOS, hacer rooting/jailbreak, etc.) sin autorización.',
        'Prohibido el uso para actividades ilegales, descarga de contenido pirata o '
        'cualquier uso que exponga a la empresa a riesgo legal o de seguridad.',
        'Prohibido conectar el equipo a redes no autorizadas cuando maneje información '
        'sensible de la empresa.',
        'El colaborador es responsable de mantener el equipo limpio, protegido '
        '(funda/maletín en el caso de laptops) y en condiciones adecuadas de operación.',
    ])

    add_heading(doc, '6. Seguridad de la información', 1)
    add_numbered(doc, [
        'Bloqueo de sesión obligatorio al alejarse del equipo.',
        'Actualizaciones de sistema operativo y antivirus corporativo no deben '
        'deshabilitarse.',
        'Información crítica del puesto debe resguardarse en las ubicaciones que indique '
        'Sistemas (no solo en el disco local del equipo).',
        'Pérdida, robo, extravío o sospecha de acceso no autorizado se reporta de '
        'inmediato a Sistemas mediante ticket urgente y, si aplica, se da parte a RH.',
    ])

    add_heading(doc, '7. Soporte y mantenimiento', 1)
    add_body(doc, 'Toda incidencia se reporta por ticket en Nexus IT, clasificado por '
                  'categoría y prioridad:')
    add_table(doc, ['Prioridad', 'Criterio', 'Tiempo de respuesta objetivo'], [
        ['Urgente', 'Equipo parado, colaborador sin poder trabajar', 'Mismo día hábil'],
        ['Alta', 'Falla que limita el trabajo pero permite operar parcialmente', '24 h'],
        ['Media', 'Incidencia que no bloquea la operación', '48–72 h'],
        ['Baja', 'Mejora o solicitud no crítica', 'Según disponibilidad de Sistemas'],
    ], widths=[1.1, 3.7, 1.7])
    add_body(doc, 'Sistemas programa mantenimientos preventivos periódicos; el colaborador '
                  'debe facilitar el acceso al equipo cuando se le solicite.')

    add_heading(doc, '8. Política de cambio y reemplazo de equipo', 1)
    add_heading(doc, '8.1 Criterios de reemplazo', 2)
    add_body(doc, 'Un equipo es candidato a reemplazo (no reparación) cuando se cumpla '
                  'cualquiera de estas condiciones:')
    add_numbered(doc, [
        'Costo de reparación mayor al 50% del costo de un equipo de reemplazo equivalente.',
        'Falla recurrente (misma falla o distinta, 3 o más veces) fuera de garantía.',
        'Daño irreparable, pérdida o robo.',
        'Obsolescencia: el equipo ya no soporta el software o los requerimientos mínimos '
        'del puesto.',
        'Fin de vida útil estimada (ver tabla 8.2), sujeto a evaluación de estado real, '
        'no solo antigüedad.',
    ])
    add_heading(doc, '8.2 Vida útil estimada por tipo de equipo', 2)
    add_body(doc, '(valores de referencia — ajustar según presupuesto y política de '
                  'depreciación de cada empresa)', color=GRAY, size=9.5)
    add_table(doc, ['Tipo de equipo', 'Vida útil estimada'], [
        ['Laptop', '4 años'],
        ['Desktop', '5 años'],
        ['Servidor', '5–7 años'],
        ['Teléfono', '3 años'],
        ['Tablet', '3 años'],
        ['Impresora', '5 años o por volumen de impresión'],
        ['Router / Switch', '5–6 años o fin de soporte del fabricante'],
    ], widths=[2.3, 4.2])
    add_heading(doc, '8.3 Procedimiento de solicitud de cambio', 2)
    add_numbered(doc, [
        'El colaborador o su jefe directo levanta un ticket categoría Hardware '
        'describiendo la falla o necesidad.',
        'Sistemas diagnostica y determina: reparación, reemplazo o continuidad sin '
        'cambios.',
        'Si procede reemplazo, Sistemas registra el dictamen y gestiona la baja del '
        'equipo anterior (estado "Baja" en Nexus IT) y el alta del nuevo, generando '
        'nueva Carta Responsiva.',
        'El equipo sustituido se resguarda o se da de baja definitiva según su estado '
        'físico.',
    ])

    add_heading(doc, '9. Responsabilidad por daño, pérdida o robo', 1)
    add_numbered(doc, [
        'Daño por negligencia comprobada (golpes, líquidos, exposición indebida, '
        'incumplimiento de esta política): el costo de reparación o reemplazo es '
        'responsabilidad del colaborador, conforme a lo firmado en su Carta Responsiva.',
        'Desgaste normal, falla de fábrica o caso fortuito (siniestro, robo con evidencia '
        'de denuncia): la empresa absorbe el costo, sujeto a revisión de Sistemas y RH.',
        'Todo caso se documenta en el ticket correspondiente antes de determinar '
        'responsabilidad.',
    ])

    add_heading(doc, '10. Baja y devolución de equipo', 1)
    add_body(doc, 'Al término de la relación laboral, cambio de puesto o cambio de empresa '
                  'dentro del grupo:')
    add_numbered(doc, [
        'El equipo se devuelve a Sistemas a más tardar en el último día laboral del '
        'colaborador. En caso de baja inmediata, despido o ausencia sin previo aviso, la '
        'devuelve su jefe directo o RH dentro de las 24 horas siguientes a la baja.',
        'Sistemas revisa el estado físico y funcional del equipo.',
        'Se elimina de forma segura la información y cuentas asociadas al colaborador '
        'saliente.',
        'Se actualiza el estado del equipo en Nexus IT (reasignable, mantenimiento o baja).',
    ])

    add_heading(doc, '11. Incumplimiento', 1)
    add_body(doc, 'El incumplimiento de esta política puede derivar en:')
    add_numbered(doc, [
        'Llamada de atención y registro en expediente.',
        'Cargo del costo de reparación o reemplazo cuando exista negligencia comprobada.',
        'Medidas disciplinarias adicionales conforme al Reglamento Interior de Trabajo '
        'vigente en cada empresa, según la gravedad del caso.',
    ])

    add_heading(doc, '12. Vigencia y actualización', 1)
    add_body(doc, 'Esta política se revisa al menos una vez al año por Sistemas y puede '
                  'actualizarse por decisión de Dirección. La versión vigente se resguarda '
                  'en Nexus IT y se comunica a todo el personal ante cualquier cambio.')

    add_heading(doc, '13. Aceptación', 1)
    add_body(doc, 'Al recibir un equipo mediante Carta Responsiva, el colaborador declara '
                  'conocer y aceptar los términos de esta Política de Uso y Cambio de '
                  'Equipo de Cómputo.')

    doc.add_paragraph().paragraph_format.space_before = Pt(20)
    sig_table = doc.add_table(rows=3, cols=2)
    labels = ['Nombre del colaborador:', 'Firma:', 'Fecha:']
    for i, label in enumerate(labels):
        sig_table.rows[i].cells[0].text = ''
        r = sig_table.rows[i].cells[0].paragraphs[0].add_run(label)
        r.font.bold = True
        r.font.size = Pt(10.5)
        sig_table.rows[i].cells[1].text = ''
        r2 = sig_table.rows[i].cells[1].paragraphs[0].add_run('_' * 40)
        r2.font.size = Pt(10.5)

    doc.save(OUT_PATH)
    print(f'Generado: {OUT_PATH}')


if __name__ == '__main__':
    build()
